from __future__ import annotations

import datetime as dt
import sqlite3
from dataclasses import dataclass
from pathlib import Path
from zoneinfo import ZoneInfo

from openpyxl import Workbook

import config
from ops_commands import (
    ACTIVE_STATUSES,
    LEDGER_STATUSES,
    OpsCommand,
    TERMINAL_STATUSES,
    is_valid_ledger_status,
)

_DISPLAY_COLS = ("task_id", "label", "status", "error", "updated_at")
_GID_COLS = (
    "task_id",
    "label",
    "status",
    "error",
    "updated_at",
    "im_delivered_at",
    "im_deliver_error",
    "session_id",
    "adb_serial",
    "buf_done_zip",
)
_EXPORT_TABLE_COLS = (
    "filename",
    "label",
    "updated_at",
    "finished_at",
)
_SHANGHAI = ZoneInfo("Asia/Shanghai")
_CHAR_BUDGET = 3500
_LIST_ROW_CAP = 20
_PROGRESS_ROW_CAP = 30
_GID_ROW_CAP = 10
_EXPORT_HARD_CAP = 50000
_LABEL_MAX = 20
_LIST_ERROR_MAX = 80
_GID_ERROR_MAX = 200


@dataclass
class LedgerQueryResult:
    ok: bool
    message: str
    file_path: Path | None = None
    row_count: int = 0
    truncated: bool = False


def _tasks_table() -> str:
    return config.TASKS_TABLE


def _connect() -> sqlite3.Connection:
    db = Path(config.TASKS_DB)
    if not db.is_file():
        raise FileNotFoundError(f"tasks db missing: {db}")
    uri = f"file:{db.resolve().as_posix()}?mode=ro"
    conn = sqlite3.connect(uri, uri=True, timeout=30.0)
    conn.row_factory = sqlite3.Row
    return conn


def to_shanghai(iso_utc: str) -> str:
    """UTC ISO (…Z or offset) → Asia/Shanghai wall time for chat display."""
    raw = (iso_utc or "").strip()
    if not raw:
        return "-"
    try:
        normalized = raw.replace("Z", "+00:00") if raw.endswith("Z") else raw
        parsed = dt.datetime.fromisoformat(normalized)
        if parsed.tzinfo is None:
            parsed = parsed.replace(tzinfo=dt.timezone.utc)
        return parsed.astimezone(_SHANGHAI).strftime("%Y-%m-%d %H:%M:%S")
    except ValueError:
        return raw


def to_shanghai_export(iso_utc: str) -> str:
    """UTC ISO → Asia/Shanghai for export table: YYYY-mm-DD: HH:mm."""
    raw = (iso_utc or "").strip()
    if not raw:
        return ""
    try:
        normalized = raw.replace("Z", "+00:00") if raw.endswith("Z") else raw
        parsed = dt.datetime.fromisoformat(normalized)
        if parsed.tzinfo is None:
            parsed = parsed.replace(tzinfo=dt.timezone.utc)
        return parsed.astimezone(_SHANGHAI).strftime("%Y-%m-%d: %H:%M")
    except ValueError:
        return raw


def _clip(text: str, max_len: int) -> str:
    s = " ".join((text or "").split())
    if len(s) <= max_len:
        return s
    if max_len <= 1:
        return s[:max_len]
    return s[: max_len - 1] + "…"


def _is_failure_terminal(status: str) -> bool:
    return status in TERMINAL_STATUSES and status != "success"


def _like_pattern(raw: str) -> str:
    """Escape LIKE wildcards then wrap with %…%."""
    escaped = (
        (raw or "")
        .replace("\\", "\\\\")
        .replace("%", "\\%")
        .replace("_", "\\_")
    )
    return f"%{escaped}%"


def _format_progress(rows: list[sqlite3.Row], total: int, *, header: str) -> str:
    lines = [header]
    for row in rows:
        label = _clip(str(row["label"] or "-"), _LABEL_MAX)
        lines.append(f"{row['task_id']}  {label}  {row['status']}")
    body = "\n".join(lines)
    if total > len(rows):
        body += (
            f"\n… showing {len(rows)}/{total}, "
            "use query gid <id> or export table"
        )
    return _fit_budget(body)


def _format_list(rows: list[sqlite3.Row], *, header: str, total: int) -> str:
    lines = [header]
    for row in rows:
        label = _clip(str(row["label"] or "-"), _LABEL_MAX)
        line = f"{row['task_id']}  {label}  {row['status']}"
        status = str(row["status"] or "")
        if _is_failure_terminal(status):
            err = _clip(str(row["error"] or ""), _LIST_ERROR_MAX)
            if err:
                line = f"{line}  {err}"
        lines.append(line)
    body = "\n".join(lines)
    if total > len(rows):
        body += (
            f"\n… showing {len(rows)}/{total}, "
            "use query gid <id> or export table"
        )
    return _fit_budget(body)


def _row_has(row: sqlite3.Row, key: str) -> bool:
    try:
        return key in row.keys()
    except Exception:
        return False


def _format_gid(rows: list[sqlite3.Row]) -> str:
    blocks: list[str] = []
    for row in rows:
        label = _clip(str(row["label"] or "-"), _LABEL_MAX)
        when = to_shanghai(str(row["updated_at"] or ""))
        head = f"{row['task_id']}  {label}  {row['status']}  {when}"
        bin_path = Path(str(row["buf_done_zip"] or "").strip())
        bin_flag = "yes" if bin_path.is_file() else "no"
        session = str(row["session_id"] or "").strip() or "-"
        adb = str(row["adb_serial"] or "").strip() or "-"
        lines = [
            head,
            f"delivered  {to_shanghai(str(row['im_delivered_at'] or ''))}",
            f"buf_done   {bin_flag}",
            f"session    {session}",
            f"adb        {adb}",
        ]
        if _row_has(row, "hotfix_has_files"):
            hf = str(row["hotfix_has_files"] or "").strip() or "-"
            src = (
                str(row["hotfix_pull_source"] or "").strip()
                if _row_has(row, "hotfix_pull_source")
                else ""
            )
            screen = (
                str(row["screen_reached"] or "").strip()
                if _row_has(row, "screen_reached")
                else ""
            )
            detail = hf
            if src:
                detail = f"{detail}/{src}"
            if screen:
                detail = f"{detail} screen={screen}"
            lines.append(f"hotfix     {detail}")
        if _row_has(row, "im_deliver_error"):
            derr = _clip(str(row["im_deliver_error"] or ""), _GID_ERROR_MAX)
            if derr:
                lines.append(f"deliver_err {derr}")
        err = _clip(str(row["error"] or ""), _GID_ERROR_MAX)
        if err:
            lines.append(err)
        blocks.append("\n".join(lines))
    return _fit_budget("\n\n".join(blocks))


def _fit_budget(text: str) -> str:
    if len(text) <= _CHAR_BUDGET:
        return text
    cut = text[: _CHAR_BUDGET - 20].rstrip()
    return cut + "\n… truncated"


def _write_export_table_xlsx(rows: list[sqlite3.Row]) -> Path:
    out_dir = Path(config.QUERY_EXPORT_DIR)
    out_dir.mkdir(parents=True, exist_ok=True)
    stamp = dt.datetime.now(_SHANGHAI).strftime("%Y%m%d_%H%M%S")
    path = out_dir / f"tasks_table_{stamp}.xlsx"
    wb = Workbook()
    ws = wb.active
    ws.title = "tasks"
    ws.append(list(_EXPORT_TABLE_COLS))
    time_cols = {"updated_at", "finished_at"}
    for row in rows:
        values = []
        for c in _EXPORT_TABLE_COLS:
            raw = row[c] if c in row.keys() else ""
            if c in time_cols:
                values.append(to_shanghai_export(str(raw or "")))
            else:
                values.append(raw)
        ws.append(values)
    wb.save(path)
    return path


def _unique_by_filename(rows: list[sqlite3.Row]) -> list[sqlite3.Row]:
    """Keep latest row per filename (rows must be ORDER BY updated_at DESC)."""
    seen: set[str] = set()
    out: list[sqlite3.Row] = []
    for row in rows:
        name = str(row["filename"] or "").strip()
        if not name or name in seen:
            continue
        seen.add(name)
        out.append(row)
    return out


def _select_active(
    conn: sqlite3.Connection,
    *,
    cols: str,
    table: str,
    sender_id: str = "",
    limit: int,
) -> tuple[int, list[sqlite3.Row]]:
    placeholders = ", ".join("?" for _ in ACTIVE_STATUSES)
    statuses = tuple(sorted(ACTIVE_STATUSES))
    if sender_id:
        total = conn.execute(
            f"SELECT COUNT(*) FROM {table} "
            f"WHERE status IN ({placeholders}) AND im_sender_id=?",
            (*statuses, sender_id),
        ).fetchone()[0]
        rows = list(
            conn.execute(
                f"SELECT {cols} FROM {table} "
                f"WHERE status IN ({placeholders}) AND im_sender_id=? "
                "ORDER BY updated_at DESC LIMIT ?",
                (*statuses, sender_id, limit),
            ).fetchall()
        )
    else:
        total = conn.execute(
            f"SELECT COUNT(*) FROM {table} WHERE status IN ({placeholders})",
            statuses,
        ).fetchone()[0]
        rows = list(
            conn.execute(
                f"SELECT {cols} FROM {table} "
                f"WHERE status IN ({placeholders}) "
                "ORDER BY updated_at DESC LIMIT ?",
                (*statuses, limit),
            ).fetchall()
        )
    return int(total), rows


def _query_gid_rows(
    conn: sqlite3.Connection, *, cols: str, table: str, gid: str
) -> list[sqlite3.Row]:
    exact = list(
        conn.execute(
            f"SELECT {cols} FROM {table} WHERE task_id=? "
            "ORDER BY updated_at DESC LIMIT 1",
            (gid,),
        ).fetchall()
    )
    if exact:
        return exact
    pattern = _like_pattern(gid)
    try:
        return list(
            conn.execute(
                f"SELECT {cols} FROM {table} "
                "WHERE label LIKE ? ESCAPE '\\' "
                "ORDER BY updated_at DESC LIMIT ?",
                (pattern, _GID_ROW_CAP),
            ).fetchall()
        )
    except sqlite3.OperationalError:
        return list(
            conn.execute(
                f"SELECT {cols} FROM {table} "
                "WHERE label LIKE ? "
                "ORDER BY updated_at DESC LIMIT ?",
                (pattern, _GID_ROW_CAP),
            ).fetchall()
        )


def _allowed_status_hint(*, include_all: bool = False) -> str:
    allowed = ", ".join(sorted(LEDGER_STATUSES))
    if include_all:
        return f"状态不合法。可用：all（全部），或 {allowed}"
    return f"状态不合法。可用：{allowed}"


def _export_usage_message() -> str:
    return (
        "用法：export table all | export table <status>\n"
        f"{_allowed_status_hint(include_all=True)}"
    )


def _query_usage_message() -> str:
    return (
        "用法：query mine | query progress | query status <status> | "
        "query gid <id> | query password\n"
        "导出请用：export table all | export table <status>"
    )


def run_ledger_query(
    cmd: OpsCommand, *, sender_id: str = ""
) -> LedgerQueryResult:
    if cmd.kind == "help":
        return LedgerQueryResult(ok=False, message=config.OPS_TEMPLATE)

    if cmd.kind == "export_usage":
        return LedgerQueryResult(ok=False, message=_export_usage_message())

    if cmd.kind == "query_usage":
        return LedgerQueryResult(ok=False, message=_query_usage_message())

    if cmd.kind == "query_password":
        password = (config.ZIP_PASSWORD or "").strip()
        if not password:
            return LedgerQueryResult(
                ok=False,
                message="ZIP_PASSWORD missing in apps/auto-extract/.env",
            )
        return LedgerQueryResult(
            ok=True,
            message=f"解压密码：{password}\n说明：将 .bin 按 zip 解压即可。",
        )

    try:
        conn = _connect()
    except FileNotFoundError as exc:
        return LedgerQueryResult(ok=False, message=str(exc))

    display_cols = ", ".join(_DISPLAY_COLS)
    gid_cols = ", ".join(_GID_COLS)
    table_cols = ", ".join(_EXPORT_TABLE_COLS)
    table = _tasks_table()
    try:
        if cmd.kind == "query_mine":
            asker = (sender_id or "").strip()
            if not asker:
                return LedgerQueryResult(
                    ok=False,
                    message="无法识别提问人，请用钉钉账号重新 @ 机器人后再 query mine",
                )
            try:
                total, rows = _select_active(
                    conn,
                    cols=display_cols,
                    table=table,
                    sender_id=asker,
                    limit=_PROGRESS_ROW_CAP,
                )
            except sqlite3.OperationalError:
                return LedgerQueryResult(
                    ok=False,
                    message="tasks 表缺少 im_sender_id，无法 query mine",
                )
            return LedgerQueryResult(
                ok=True,
                message=_format_progress(
                    rows,
                    total,
                    header=f"mine: {total} active (sender={asker})",
                ),
                row_count=len(rows),
                truncated=total > len(rows),
            )

        if cmd.kind == "query_progress":
            total, rows = _select_active(
                conn,
                cols=display_cols,
                table=table,
                limit=_PROGRESS_ROW_CAP,
            )
            return LedgerQueryResult(
                ok=True,
                message=_format_progress(
                    rows, total, header=f"progress: {total} active"
                ),
                row_count=len(rows),
                truncated=total > len(rows),
            )

        if cmd.kind == "query_status":
            status = (cmd.arg or "").strip()
            if not is_valid_ledger_status(status):
                return LedgerQueryResult(
                    ok=False,
                    message=_allowed_status_hint(),
                )
            total = conn.execute(
                f"SELECT COUNT(*) FROM {table} WHERE status=?", (status,)
            ).fetchone()[0]
            rows = list(
                conn.execute(
                    f"SELECT {display_cols} FROM {table} WHERE status=? "
                    "ORDER BY updated_at DESC LIMIT ?",
                    (status, _LIST_ROW_CAP),
                ).fetchall()
            )
            return LedgerQueryResult(
                ok=True,
                message=_format_list(
                    rows,
                    header=f"status={status}  {len(rows)} shown",
                    total=int(total),
                ),
                row_count=len(rows),
                truncated=int(total) > len(rows),
            )

        if cmd.kind == "query_gid":
            gid = (cmd.arg or "").strip()
            if not gid:
                return LedgerQueryResult(
                    ok=False,
                    message="用法：query gid <task_id|label>",
                )
            try:
                rows = _query_gid_rows(
                    conn, cols=gid_cols, table=table, gid=gid
                )
            except sqlite3.OperationalError:
                # Schema without im_deliver_error: drop that column.
                slim = ", ".join(c for c in _GID_COLS if c != "im_deliver_error")
                rows = _query_gid_rows(
                    conn, cols=slim, table=table, gid=gid
                )
            if not rows:
                return LedgerQueryResult(
                    ok=True, message=f"not found: {gid}", row_count=0
                )
            msg = _format_gid(rows)
            if len(rows) >= _GID_ROW_CAP:
                msg += (
                    f"\n… capped at {_GID_ROW_CAP}, "
                    "refine query or use export table"
                )
            return LedgerQueryResult(
                ok=True, message=msg, row_count=len(rows)
            )

        if cmd.kind == "export_table":
            scope = (cmd.arg or "").strip().lower()
            if not scope:
                return LedgerQueryResult(
                    ok=False,
                    message=_export_usage_message(),
                )
            if scope != "all" and not is_valid_ledger_status(scope):
                return LedgerQueryResult(
                    ok=False,
                    message=_allowed_status_hint(include_all=True),
                )
            # Over-fetch then unique by filename (latest updated_at wins).
            fetch_cap = _EXPORT_HARD_CAP * 3
            if scope == "all":
                rows = list(
                    conn.execute(
                        f"SELECT {table_cols} FROM {table} "
                        "WHERE TRIM(COALESCE(filename,'')) != '' "
                        "ORDER BY updated_at DESC LIMIT ?",
                        (fetch_cap,),
                    ).fetchall()
                )
            else:
                rows = list(
                    conn.execute(
                        f"SELECT {table_cols} FROM {table} "
                        "WHERE status=? AND TRIM(COALESCE(filename,'')) != '' "
                        "ORDER BY updated_at DESC LIMIT ?",
                        (scope, fetch_cap),
                    ).fetchall()
                )
            unique = _unique_by_filename(rows)
            truncated = len(unique) > _EXPORT_HARD_CAP
            if truncated:
                unique = unique[:_EXPORT_HARD_CAP]
            path = _write_export_table_xlsx(unique)
            scope_label = "all" if scope == "all" else f"status={scope}"
            msg = (
                f"export table {scope_label}: {len(unique)} unique filenames "
                f"xlsx (Asia/Shanghai YYYY-mm-DD: HH:mm)"
            )
            if truncated:
                msg += f" truncated=true cap={_EXPORT_HARD_CAP}"
            return LedgerQueryResult(
                ok=True,
                message=msg,
                file_path=path,
                row_count=len(unique),
                truncated=truncated,
            )

        return LedgerQueryResult(ok=False, message=_query_usage_message())
    finally:
        conn.close()
